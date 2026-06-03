import csv
import io
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook


REQUIRED_COLUMNS = {'account_code', 'account_name', 'debit', 'credit'}
OPTIONAL_COLUMNS = {'period', 'description'}


def _normalize_columns(columns: list[str]) -> dict[str, str]:
    """Map various column names to standardized keys."""
    mapping = {}
    col_map = {
        'account_code': ['account_code', '科目代码', '科目编码', '科目编号', 'code'],
        'account_name': ['account_name', '科目名称', '科目名', 'name'],
        'debit': ['debit', '借方', '借方金额', 'debit_amount'],
        'credit': ['credit', '贷方', '贷方金额', 'credit_amount'],
        'period': ['period', '期间', '会计期间', '月份'],
        'description': ['description', '摘要', '说明', '备注'],
    }
    lower_cols = [c.strip().lower().replace(' ', '_') for c in columns]
    for std_key, variants in col_map.items():
        for variant in variants:
            if variant in lower_cols:
                idx = lower_cols.index(variant)
                mapping[std_key] = columns[idx]
                break
    return mapping


def parse_ledger_file(file_obj, file_name: str):
    """
    Parse CSV or Excel file into a list of row dicts.
    Returns {
        'rows': [...],
        'errors': [...],
        'summary': {...}
    }
    """
    ext = file_name.split('.')[-1].lower()
    try:
        if ext == 'csv':
            content = file_obj.read().decode('utf-8-sig')
            reader = csv.DictReader(io.StringIO(content))
            raw_rows = list(reader)
            columns = reader.fieldnames or []
        elif ext in ['xlsx', 'xls']:
            file_obj.seek(0)
            wb = load_workbook(file_obj, data_only=True)
            ws = wb.active
            columns = [str(cell.value or '') for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            raw_rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                raw_rows.append(dict(zip(columns, row)))
        else:
            return {'rows': [], 'errors': [f'不支持的文件格式: {ext}'], 'summary': {}}
    except Exception as e:
        return {'rows': [], 'errors': [f'文件解析失败: {str(e)}'], 'summary': {}}

    col_mapping = _normalize_columns(columns)
    missing = REQUIRED_COLUMNS - set(col_mapping.keys())
    if missing:
        return {
            'rows': [],
            'errors': [f'缺少必需列: {", ".join(missing)}。检测到列: {columns}'],
            'summary': {}
        }

    rows = []
    errors = []
    total_debit = Decimal('0')
    total_credit = Decimal('0')

    for idx, row in enumerate(raw_rows):
        row_num = idx + 2
        account_code = str(row.get(col_mapping['account_code'], '') or '').strip()
        account_name = str(row.get(col_mapping['account_name'], '') or '').strip()

        if not account_code or not account_name:
            errors.append({'row': row_num, 'error': '科目代码或名称为空'})
            continue

        try:
            debit_val = str(row.get(col_mapping['debit'], '0') or '0').strip().replace(',', '')
            credit_val = str(row.get(col_mapping['credit'], '0') or '0').strip().replace(',', '')
            debit = Decimal(debit_val or '0')
            credit = Decimal(credit_val or '0')
        except InvalidOperation:
            errors.append({'row': row_num, 'error': '借贷方金额格式错误'})
            continue

        period_col = col_mapping.get('period')
        period = str(row.get(period_col, '') or '').strip() if period_col else 'default'

        desc_col = col_mapping.get('description')
        description = str(row.get(desc_col, '') or '').strip() if desc_col else ''

        rows.append({
            'account_code': account_code,
            'account_name': account_name,
            'debit': debit,
            'credit': credit,
            'period': period,
            'description': description,
        })
        total_debit += debit
        total_credit += credit

    summary = {
        'total_rows': len(raw_rows),
        'parsed_rows': len(rows),
        'error_rows': len(errors),
        'total_debit': float(total_debit),
        'total_credit': float(total_credit),
    }

    return {'rows': rows, 'errors': errors, 'summary': summary}
