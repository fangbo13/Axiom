import csv
import io
import random
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook


TB_COLUMN_MAP = {
    'account_code': ['account_code', '科目代码', '科目编码', '科目编号', 'code', '科目号'],
    'account_name': ['account_name', '科目名称', '科目名', 'name'],
    'opening_debit': ['opening_debit', '期初借方', '期初借方余额', '年初借方', '期初借'],
    'opening_credit': ['opening_credit', '期初贷方', '期初贷方余额', '年初贷方', '期初贷'],
    'period_debit': ['period_debit', '本期借方', '本期借方发生额', '借方发生额', '本期借'],
    'period_credit': ['period_credit', '本期贷方', '本期贷方发生额', '贷方发生额', '本期贷'],
    'closing_debit': ['closing_debit', '期末借方', '期末借方余额', '年末借方', '期末借'],
    'closing_credit': ['closing_credit', '期末贷方', '期末贷方余额', '年末贷方', '期末贷'],
    'direction': ['direction', '方向', '余额方向', '借贷方向'],
}

JE_COLUMN_MAP = {
    'voucher_date': ['voucher_date', '凭证日期', '日期', '记账日期'],
    'voucher_no': ['voucher_no', '凭证号', '凭证编号', '记字号'],
    'line_no': ['line_no', '行号', '分录号'],
    'abstract': ['abstract', '摘要', '说明'],
    'account_code': ['account_code', '科目代码', '科目编码', '科目编号', 'code'],
    'account_name': ['account_name', '科目名称', '科目名', 'name'],
    'debit': ['debit', '借方', '借方金额', 'debit_amount'],
    'credit': ['credit', '贷方', '贷方金额', 'credit_amount'],
}


def _normalize_columns(columns: list[str], column_map: dict):
    """Map various column names to standardized keys."""
    mapping = {}
    lower_cols = [c.strip().lower().replace(' ', '_') for c in columns]
    for std_key, variants in column_map.items():
        for variant in variants:
            if variant in lower_cols:
                idx = lower_cols.index(variant)
                mapping[std_key] = columns[idx]
                break
    return mapping


def parse_file_to_raw_rows(file_obj, file_name: str):
    """
    Parse CSV or Excel file into raw row dicts.
    Returns {
        'columns': [...],
        'rows': [...],
        'errors': [],
    }
    """
    ext = file_name.split('.')[-1].lower()
    try:
        if ext == 'csv':
            content = file_obj.read().decode('utf-8-sig')
            reader = csv.DictReader(io.StringIO(content))
            raw_rows = list(reader)
            columns = reader.fieldnames or []
        elif ext in ['xlsx', 'xls']:
            file_obj.seek(0)
            wb = load_workbook(file_obj, data_only=True)
            ws = wb.active
            columns = [str(cell.value or '') for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            raw_rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                raw_rows.append(dict(zip(columns, row)))
        else:
            return {'columns': [], 'rows': [], 'errors': [f'不支持的文件格式: {ext}']}
    except Exception as e:
        return {'columns': [], 'rows': [], 'errors': [f'文件解析失败: {str(e)}']}

    return {'columns': columns, 'rows': raw_rows, 'errors': []}


def guess_column_mapping(columns: list[str], import_type: str):
    """Guess column mapping based on import type."""
    col_map = TB_COLUMN_MAP if import_type == 'tb' else JE_COLUMN_MAP
    return _normalize_columns(columns, col_map)


def get_sample_rows(raw_rows: list[dict], count: int = 100, random_samples: int = 10):
    """Get first N rows plus random samples from middle/end."""
    first_n = raw_rows[:count]
    remaining = raw_rows[count:]
    samples = []
    if remaining and random_samples > 0:
        samples = random.sample(remaining, min(random_samples, len(remaining)))
    return first_n, samples


def apply_mapping(raw_rows: list[dict], column_mapping: dict, import_type: str):
    """
    Transform raw rows into structured data using confirmed column mapping.
    Returns {
        'rows': [...],
        'errors': [{'row': int, 'error': str}],
    }
    """
    rows = []
    errors = []

    for idx, raw in enumerate(raw_rows):
        row_num = idx + 2
        mapped = {}

        for std_key, source_col in column_mapping.items():
            mapped[std_key] = raw.get(source_col, '')

        if import_type == 'tb':
            account_code = str(mapped.get('account_code', '') or '').strip()
            account_name = str(mapped.get('account_name', '') or '').strip()

            if not account_code:
                errors.append({'row': row_num, 'error': '科目代码为空'})
                continue

            try:
                opening_debit = _to_decimal(mapped.get('opening_debit'))
                opening_credit = _to_decimal(mapped.get('opening_credit'))
                period_debit = _to_decimal(mapped.get('period_debit'))
                period_credit = _to_decimal(mapped.get('period_credit'))
                closing_debit = _to_decimal(mapped.get('closing_debit'))
                closing_credit = _to_decimal(mapped.get('closing_credit'))
            except (InvalidOperation, ValueError):
                errors.append({'row': row_num, 'error': '金额格式错误'})
                continue

            direction = str(mapped.get('direction', '') or '').strip()
            if direction in ['借', '借方', 'debit', 'dr']:
                direction = 'debit'
            elif direction in ['贷', '贷方', 'credit', 'cr']:
                direction = 'credit'
            else:
                direction = ''

            rows.append({
                'account_code': account_code,
                'account_name': account_name,
                'opening_debit': opening_debit,
                'opening_credit': opening_credit,
                'period_debit': period_debit,
                'period_credit': period_credit,
                'closing_debit': closing_debit,
                'closing_credit': closing_credit,
                'direction': direction,
            })
        else:
            # JE type
            account_code = str(mapped.get('account_code', '') or '').strip()
            if not account_code:
                errors.append({'row': row_num, 'error': '科目代码为空'})
                continue

            try:
                debit = _to_decimal(mapped.get('debit'))
                credit = _to_decimal(mapped.get('credit'))
            except (InvalidOperation, ValueError):
                errors.append({'row': row_num, 'error': '金额格式错误'})
                continue

            voucher_date = mapped.get('voucher_date', '') or ''
            voucher_no = str(mapped.get('voucher_no', '') or '').strip()
            line_no = _to_int(mapped.get('line_no'), default=1)
            abstract = str(mapped.get('abstract', '') or '').strip()
            account_name = str(mapped.get('account_name', '') or '').strip()

            rows.append({
                'voucher_date': voucher_date,
                'voucher_no': voucher_no,
                'line_no': line_no,
                'abstract': abstract,
                'account_code': account_code,
                'account_name': account_name,
                'debit': debit,
                'credit': credit,
            })

    return {'rows': rows, 'errors': errors}


def _to_decimal(value, default=Decimal('0')):
    if value is None or value == '':
        return default
    val = str(value).strip().replace(',', '').replace('，', '')
    return Decimal(val or '0')


def _to_int(value, default=0):
    if value is None or value == '':
        return default
    try:
        return int(str(value).strip())
    except (ValueError, TypeError):
        return default
