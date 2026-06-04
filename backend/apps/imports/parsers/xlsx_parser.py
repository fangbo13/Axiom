from openpyxl import load_workbook
from .base import BaseParser, ParseResult
from .header_detector import HeaderDetector
from .merged_cell_resolver import MergedCellResolver


class XlsxParser(BaseParser):

    @classmethod
    def can_handle(cls, file_name: str) -> bool:
        return file_name.lower().endswith(('.xlsx', '.xls'))

    def parse(self, file_obj, file_name: str, options: dict) -> ParseResult:
        result = ParseResult()
        header_row_hint = options.get('header_row_hint')
        max_scan = options.get('max_scan', 30)

        try:
            file_obj.seek(0)
            wb = load_workbook(file_obj, data_only=True)
            ws = wb.active
        except Exception as e:
            result.errors.append(f'Excel解析失败: {str(e)}')
            return result

        # Read all rows into a 2D list
        rows_2d = []
        max_cols = 0
        for row in ws.iter_rows(values_only=True):
            row_list = list(row)
            rows_2d.append(row_list)
            max_cols = max(max_cols, len(row_list))

        if not rows_2d:
            result.errors.append('Excel文件为空')
            return result

        # Resolve merged cells
        merged_info = MergedCellResolver.resolve(ws, rows_2d)
        result.merged_cells_info = merged_info

        # Detect header row
        detector = HeaderDetector()
        detection = detector.detect(rows_2d, max_scan=max_scan, hint=header_row_hint)
        result.header_row = detection.row
        result.header_confidence = detection.confidence
        result.is_multiline_header = detection.is_multiline

        # Extract columns
        header_idx = detection.row
        if detection.is_multiline and header_idx + 1 < len(rows_2d):
            # Merge two-level headers
            parent_row = rows_2d[header_idx]
            child_row = rows_2d[header_idx + 1]
            columns = []
            current_parent = ''
            for i in range(max(len(parent_row), len(child_row))):
                parent = str(parent_row[i] or '').strip() if i < len(parent_row) else ''
                child = str(child_row[i] or '').strip() if i < len(child_row) else ''
                if parent:
                    current_parent = parent
                if child:
                    if current_parent and current_parent != child:
                        columns.append(f"{current_parent}_{child}")
                    else:
                        columns.append(child)
                else:
                    columns.append(current_parent or f"列{i+1}")
            data_start = header_idx + 2
        else:
            columns = [str(c or '').strip() for c in rows_2d[header_idx]]
            data_start = header_idx + 1

        result.columns = columns

        # Extract data rows
        for row in rows_2d[data_start:]:
            row_dict = {}
            has_value = False
            for i, col_name in enumerate(columns):
                val = row[i] if i < len(row) else None
                row_dict[col_name] = val
                if val is not None and str(val).strip() != '':
                    has_value = True
            if has_value:
                result.rows.append(row_dict)

        return result
