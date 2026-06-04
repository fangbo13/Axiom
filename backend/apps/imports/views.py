import os
from django.conf import settings
from django.db import transaction
from django.http import FileResponse
from rest_framework import viewsets, status, filters
from rest_framework.decorators import action
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from rest_framework.response import Response
from common.pagination import StandardResultsSetPagination
from common.permissions import IsProjectMember
from apps.projects.models import Project
from apps.adjustments.services import recalculate_trial_balance
from apps.ledger.models import Currency
from .models import (
    ImportBatch, UnauditedTB, UnauditedJE, ImportErrorRow,
    MappingTemplate, ImportLog,
)
from .serializers import (
    ImportBatchSerializer, ImportBatchListSerializer,
    UnauditedTBSerializer, UnauditedJESerializer,
    ImportErrorRowSerializer, PreviewResponseSerializer,
    ValidationResponseSerializer, ImportCommitSerializer,
    ColumnMappingSerializer, MappingTemplateSerializer,
    ImportLogSerializer,
)
from .parsers import parse_file, ColumnMapper, get_sample_rows, apply_mapping
from .validators import ImportValidator
from .services import get_next_version, activate_batch, commit_batch
from .async_tasks import submit_import_job, get_job_status


class ImportPipelineViewSet(viewsets.ViewSet):
    permission_classes = [IsProjectMember]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    @action(detail=False, methods=['post'])
    def upload(self, request, project_id=None):
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'detail': '未提供文件'}, status=status.HTTP_400_BAD_REQUEST)

        import_type = request.data.get('import_type', 'tb')
        overwrite_mode = request.data.get('overwrite_mode', 'append')
        period = request.data.get('period', '')

        ext = file_obj.name.split('.')[-1].lower()
        if ext not in ['csv', 'xlsx', 'xls', 'tsv', 'txt']:
            return Response({'detail': '只支持 CSV、TSV、TXT 或 Excel 文件'}, status=status.HTTP_400_BAD_REQUEST)

        project = Project.objects.get(id=project_id)
        version = get_next_version(project, period, import_type)

        # Check existing batches for overwrite prompt
        existing_batches = ImportBatch.objects.filter(
            project=project, period=period, import_type=import_type, status='committed'
        ).order_by('-version')
        existing_info = [
            {'id': b.id, 'version': b.version, 'file_name': b.file_name, 'is_active': b.is_active}
            for b in existing_batches
        ]

        batch = ImportBatch.objects.create(
            project=project,
            uploaded_by=request.user,
            file_name=file_obj.name,
            import_type=import_type,
            status='uploaded',
            overwrite_mode=overwrite_mode,
            period=period,
            version=version,
        )

        # Save file to disk for later re-reading
        batch_dir = os.path.join(settings.MEDIA_ROOT, 'imports', f'project_{project_id}', f'batch_{batch.id}')
        os.makedirs(batch_dir, exist_ok=True)
        file_path = os.path.join(batch_dir, file_obj.name)
        with open(file_path, 'wb+') as dest:
            for chunk in file_obj.chunks():
                dest.write(chunk)

        # Parse preview
        header_row_hint = request.data.get('header_row_hint')
        if header_row_hint is not None:
            try:
                header_row_hint = int(header_row_hint)
            except (ValueError, TypeError):
                header_row_hint = None
        try:
            with open(file_path, 'rb') as f:
                result = parse_file(f, file_obj.name, options={'header_row_hint': header_row_hint})
        except Exception as e:
            batch.status = 'failed'
            batch.save()
            return Response({'detail': f'文件解析失败: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        if result.errors:
            batch.status = 'failed'
            batch.save()
            return Response({'detail': '文件解析失败', 'errors': result.errors}, status=status.HTTP_400_BAD_REQUEST)

        columns = result.columns
        raw_rows = result.rows
        guessed_mapping = ColumnMapper.guess_mapping(columns, import_type)
        first_rows, sample_rows = get_sample_rows(raw_rows)

        # Detect single currency
        detected_currencies = result.detected_currencies or []
        if not detected_currencies and 'currency' in guessed_mapping:
            curr_col = guessed_mapping['currency']
            currencies = set(str(r.get(curr_col, '')).strip() for r in raw_rows if r.get(curr_col))
            detected_currencies = list(currencies)

        batch.total_rows = len(raw_rows)
        batch.metadata = {
            'file_path': file_path,
            'columns': columns,
            'guessed_mapping': guessed_mapping,
            'sample_row_indices': [i + len(first_rows) + 2 for i in range(len(sample_rows))],
            'detected_header_row': result.header_row,
            'header_confidence': result.header_confidence,
            'is_multiline_header': result.is_multiline_header,
            'merged_cells_info': result.merged_cells_info,
            'detected_currencies': detected_currencies,
            'parser_engine': ext,
        }
        batch.save()

        # Log upload
        ImportLog.objects.create(
            project=project,
            import_batch=batch,
            action='upload',
            performed_by=request.user,
            details={'file_name': file_obj.name, 'total_rows': len(raw_rows), 'period': period}
        )

        return Response({
            'batch_id': batch.id,
            'file_name': batch.file_name,
            'import_type': batch.import_type,
            'detected_columns': columns,
            'guessed_mapping': guessed_mapping,
            'total_rows': batch.total_rows,
            'first_rows': first_rows,
            'sample_rows': sample_rows,
            'detected_header_row': result.header_row,
            'header_confidence': result.header_confidence,
            'is_multiline_header': result.is_multiline_header,
            'detected_currencies': detected_currencies,
            'existing_batches': existing_info,
        })

    @action(detail=False, methods=['post'])
    def preview(self, request, project_id=None):
        batch_id = request.data.get('batch_id')
        header_row = request.data.get('header_row')
        if header_row is not None:
            try:
                header_row = int(header_row)
            except (ValueError, TypeError):
                header_row = None
        column_mapping = request.data.get('column_mapping', {})

        try:
            batch = ImportBatch.objects.get(id=batch_id, project_id=project_id)
        except ImportBatch.DoesNotExist:
            return Response({'detail': '批次不存在'}, status=status.HTTP_404_NOT_FOUND)

        file_path = batch.metadata.get('file_path')
        if not file_path or not os.path.exists(file_path):
            return Response({'detail': '文件已丢失，请重新上传'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            with open(file_path, 'rb') as f:
                result = parse_file(f, batch.file_name, options={'header_row_hint': header_row})
        except Exception as e:
            return Response({'detail': f'文件解析失败: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        if result.errors:
            return Response({'detail': '文件解析失败', 'errors': result.errors}, status=status.HTTP_400_BAD_REQUEST)

        parsed = apply_mapping(result.rows[:200], column_mapping, batch.import_type)

        return Response({
            'batch_id': batch.id,
            'header_row': result.header_row,
            'header_confidence': result.header_confidence,
            'is_multiline_header': result.is_multiline_header,
            'preview_rows': parsed['rows'],
            'preview_errors': parsed['errors'],
        })

    @action(detail=False, methods=['post'])
    def column_map(self, request, project_id=None):
        batch_id = request.data.get('batch_id')
        column_mapping = request.data.get('column_mapping', {})
        save_template = request.data.get('save_template', False)
        template_name = request.data.get('template_name', '')

        try:
            batch = ImportBatch.objects.get(id=batch_id, project_id=project_id)
        except ImportBatch.DoesNotExist:
            return Response({'detail': '批次不存在'}, status=status.HTTP_404_NOT_FOUND)

        batch.metadata['column_mapping'] = column_mapping
        batch.metadata['column_mapping_source'] = 'manual'
        batch.status = 'preview'
        batch.save()

        # Auto-save mapping template
        if column_mapping:
            header_signature = ColumnMapper.compute_header_signature(batch.metadata.get('columns', []))
            existing_auto = MappingTemplate.objects.filter(
                project=batch.project,
                import_type=batch.import_type,
                is_auto_saved=True,
            ).first()
            if existing_auto and existing_auto.header_signature.get('hash') == header_signature.get('hash'):
                existing_auto.name = f"自动保存-{batch.import_type}-{batch.metadata.get('parser_engine', 'unknown')}"
                existing_auto.file_extensions = [batch.metadata.get('parser_engine', '')]
                existing_auto.header_signature = header_signature
                existing_auto.column_mapping = column_mapping
                existing_auto.save()
            else:
                MappingTemplate.objects.create(
                    project=batch.project,
                    name=f"自动保存-{batch.import_type}-{batch.metadata.get('parser_engine', 'unknown')}",
                    import_type=batch.import_type,
                    file_extensions=[batch.metadata.get('parser_engine', '')],
                    header_signature=header_signature,
                    column_mapping=column_mapping,
                    is_auto_saved=True,
                )

        if save_template and template_name:
            header_signature = ColumnMapper.compute_header_signature(batch.metadata.get('columns', []))
            MappingTemplate.objects.create(
                project=batch.project,
                name=template_name,
                import_type=batch.import_type,
                file_extensions=[batch.metadata.get('parser_engine', '')],
                header_signature=header_signature,
                column_mapping=column_mapping,
                is_auto_saved=False,
            )

        ImportLog.objects.create(
            project=batch.project,
            import_batch=batch,
            action='map',
            performed_by=request.user,
            details={'column_mapping': column_mapping}
        )

        return Response({'batch_id': batch.id, 'status': batch.status, 'column_mapping': column_mapping})

    @action(detail=False, methods=['post'])
    def validate(self, request, project_id=None):
        batch_id = request.data.get('batch_id')
        column_mapping = request.data.get('column_mapping')

        try:
            batch = ImportBatch.objects.get(id=batch_id, project_id=project_id)
        except ImportBatch.DoesNotExist:
            return Response({'detail': '批次不存在'}, status=status.HTTP_404_NOT_FOUND)

        file_path = batch.metadata.get('file_path')
        if not file_path or not os.path.exists(file_path):
            return Response({'detail': '文件已丢失，请重新上传'}, status=status.HTTP_400_BAD_REQUEST)

        mapping = column_mapping or batch.metadata.get('column_mapping') or batch.metadata.get('guessed_mapping')

        try:
            with open(file_path, 'rb') as f:
                result = parse_file(f, batch.file_name)
        except Exception as e:
            return Response({'detail': f'文件解析失败: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        parsed = apply_mapping(result.rows, mapping, batch.import_type)

        # Run validation
        validator = ImportValidator(batch.import_type, parsed['rows'], mapping, project_id=project_id)
        checks, validation_errors, validation_warnings = validator.validate()

        # Store validation errors and warnings
        ImportErrorRow.objects.filter(import_batch=batch).delete()
        for err in validation_errors:
            ImportErrorRow.objects.create(
                import_batch=batch,
                row_number=err.get('row_number') or 0,
                raw_data=err.get('raw_data', {}),
                error_type=err['error_type'],
                error_message=err['error_message'],
                severity=err.get('severity', 'error'),
            )
        for warn in validation_warnings:
            ImportErrorRow.objects.create(
                import_batch=batch,
                row_number=warn.get('row_number') or 0,
                raw_data=warn.get('raw_data', {}),
                error_type=warn['error_type'],
                error_message=warn['error_message'],
                severity=warn.get('severity', 'warning'),
            )

        batch.status = 'validated'
        batch.parsed_rows = len(parsed['rows'])
        batch.error_rows_count = len(validation_errors)
        batch.warning_rows_count = len(validation_warnings)
        batch.validation_summary = {
            'checks': checks,
            'total_debit': str(sum(r.get('debit', 0) for r in parsed['rows'])),
            'total_credit': str(sum(r.get('credit', 0) for r in parsed['rows'])),
        }
        if batch.import_type == 'tb':
            batch.validation_summary['total_opening_debit'] = str(sum(r.get('opening_debit', 0) for r in parsed['rows']))
            batch.validation_summary['total_opening_credit'] = str(sum(r.get('opening_credit', 0) for r in parsed['rows']))
            batch.validation_summary['total_closing_debit'] = str(sum(r.get('closing_debit', 0) for r in parsed['rows']))
            batch.validation_summary['total_closing_credit'] = str(sum(r.get('closing_credit', 0) for r in parsed['rows']))
        batch.save()

        ImportLog.objects.create(
            project=batch.project,
            import_batch=batch,
            action='validate',
            performed_by=request.user,
            details={'error_count': len(validation_errors), 'warning_count': len(validation_warnings)}
        )

        return Response({
            'batch_id': batch.id,
            'is_valid': len(validation_errors) == 0,
            'summary': {
                'total_rows': batch.total_rows,
                'parsed_rows': batch.parsed_rows,
                'error_rows': batch.error_rows_count,
                'warning_rows': batch.warning_rows_count,
            },
            'checks': checks,
            'errors': validation_errors,
            'warnings': validation_warnings,
        })

    @action(detail=False, methods=['post'])
    def commit(self, request, project_id=None):
        serializer = ImportCommitSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        batch_id = serializer.validated_data['batch_id']
        set_active = serializer.validated_data.get('set_active', True)

        try:
            batch = ImportBatch.objects.get(id=batch_id, project_id=project_id)
        except ImportBatch.DoesNotExist:
            return Response({'detail': '批次不存在'}, status=status.HTTP_404_NOT_FOUND)

        if batch.status != 'validated':
            return Response({'detail': '批次尚未校验或校验失败'}, status=status.HTTP_400_BAD_REQUEST)

        # Check unresolved errors
        unresolved = batch.import_errors.filter(is_resolved=False, severity='error').count()
        if unresolved > 0:
            return Response({
                'detail': f'存在 {unresolved} 条未修正的错误，请先修正或确认忽略',
                'unresolved_errors': unresolved,
            }, status=status.HTTP_400_BAD_REQUEST)

        file_path = batch.metadata.get('file_path')
        mapping = batch.metadata.get('column_mapping') or batch.metadata.get('guessed_mapping')

        try:
            with open(file_path, 'rb') as f:
                result = parse_file(f, batch.file_name)
        except Exception as e:
            return Response({'detail': f'文件解析失败: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        parsed = apply_mapping(result.rows, mapping, batch.import_type)

        # Handle overwrite mode
        if batch.overwrite_mode == 'replace':
            ImportBatch.objects.filter(
                project_id=project_id, period=batch.period, import_type=batch.import_type
            ).exclude(id=batch.id).update(is_active=False)

        # Async for large files
        if len(parsed['rows']) > 5000:
            batch.status = 'processing'
            batch.save(update_fields=['status'])
            job_id = submit_import_job(batch, parsed['rows'])
            return Response({
                'batch_id': batch.id,
                'status': 'processing',
                'job_id': job_id,
                'message': '大数据量导入已提交异步处理',
            })

        commit_batch(batch, parsed['rows'])

        if set_active:
            activate_batch(batch)
            recalculate_trial_balance(project_id)

        ImportLog.objects.create(
            project=batch.project,
            import_batch=batch,
            action='commit',
            performed_by=request.user,
            details={'records_created': batch.parsed_rows, 'set_active': set_active}
        )

        return Response({
            'batch_id': batch.id,
            'status': batch.status,
            'records_created': batch.parsed_rows,
            'is_active': batch.is_active,
            'trial_balance_recalculated': set_active,
        })

    @action(detail=False, methods=['post'])
    def export_errors(self, request, project_id=None):
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        import tempfile

        batch_id = request.data.get('batch_id')
        try:
            batch = ImportBatch.objects.get(id=batch_id, project_id=project_id)
        except ImportBatch.DoesNotExist:
            return Response({'detail': '批次不存在'}, status=status.HTTP_404_NOT_FOUND)

        error_rows = batch.import_errors.all()
        if not error_rows.exists():
            return Response({'detail': '没有错误行可导出'}, status=status.HTTP_400_BAD_REQUEST)

        wb = Workbook()
        ws = wb.active
        ws.title = '错误行'

        headers = ['行号', '错误类型', '严重程度', '错误信息', '原始数据']
        ws.append(headers)

        for err in error_rows:
            ws.append([
                err.row_number,
                err.get_error_type_display(),
                err.get_severity_display(),
                err.error_message,
                str(err.raw_data),
            ])

        # Auto-adjust column widths
        for col_idx in range(1, len(headers) + 1):
            max_length = 0
            col_letter = get_column_letter(col_idx)
            for cell in ws[col_letter]:
                try:
                    max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 2, 60)

        temp_file = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp_file.name)
        temp_file.close()

        return FileResponse(open(temp_file.name, 'rb'), as_attachment=True, filename=f'errors_batch_{batch.id}.xlsx')

    @action(detail=False, methods=['get'])
    def job_status(self, request, project_id=None):
        job_id = request.query_params.get('job_id')
        if not job_id:
            return Response({'detail': '缺少job_id'}, status=status.HTTP_400_BAD_REQUEST)
        job = get_job_status(job_id)
        return Response(job)


class ImportBatchViewSet(viewsets.ModelViewSet):
    permission_classes = [IsProjectMember]
    pagination_class = StandardResultsSetPagination
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['file_name', 'period']
    ordering_fields = ['created_at', 'version', 'period']

    def get_queryset(self):
        queryset = ImportBatch.objects.filter(project_id=self.kwargs['project_id'])
        status_filter = self.request.query_params.get('status')
        period_filter = self.request.query_params.get('period')
        type_filter = self.request.query_params.get('import_type')
        if status_filter:
            queryset = queryset.filter(status=status_filter)
        if period_filter:
            queryset = queryset.filter(period=period_filter)
        if type_filter:
            queryset = queryset.filter(import_type=type_filter)
        return queryset

    def get_serializer_class(self):
        if self.action == 'list':
            return ImportBatchListSerializer
        return ImportBatchSerializer

    def perform_create(self, serializer):
        project = Project.objects.get(id=self.kwargs['project_id'])
        serializer.save(project=project, uploaded_by=self.request.user)

    def perform_destroy(self, instance):
        # Log deletion
        ImportLog.objects.create(
            project=instance.project,
            import_batch=instance,
            action='delete',
            performed_by=self.request.user,
            details={'file_name': instance.file_name, 'period': instance.period}
        )
        # Delete physical file
        file_path = instance.metadata.get('file_path')
        if file_path and os.path.exists(file_path):
            os.remove(file_path)
        super().perform_destroy(instance)

    @action(detail=True, methods=['post'])
    def set_active(self, request, project_id=None, pk=None):
        batch = self.get_object()
        activate_batch(batch)
        recalculate_trial_balance(project_id)
        ImportLog.objects.create(
            project=batch.project,
            import_batch=batch,
            action='activate',
            performed_by=request.user,
            details={'version': batch.version}
        )
        return Response({'detail': '已设为当前激活版本', 'batch_id': batch.id})

    @action(detail=True, methods=['get'])
    def error_rows(self, request, project_id=None, pk=None):
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
        import tempfile

        batch = self.get_object()
        rows = batch.import_errors.all()
        format_param = request.query_params.get('format')
        if format_param == 'excel':
            if not rows.exists():
                return Response({'detail': '没有错误行可导出'}, status=status.HTTP_400_BAD_REQUEST)
            wb = Workbook()
            ws = wb.active
            ws.title = '错误行'
            headers = ['行号', '错误类型', '严重程度', '错误信息', '原始数据']
            ws.append(headers)
            for err in rows:
                ws.append([
                    err.row_number,
                    err.get_error_type_display(),
                    err.get_severity_display(),
                    err.error_message,
                    str(err.raw_data),
                ])
            for col_idx in range(1, len(headers) + 1):
                max_length = 0
                col_letter = get_column_letter(col_idx)
                for cell in ws[col_letter]:
                    try:
                        max_length = max(max_length, len(str(cell.value)))
                    except:
                        pass
                ws.column_dimensions[col_letter].width = min(max_length + 2, 60)
            temp_file = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
            wb.save(temp_file.name)
            temp_file.close()
            return FileResponse(open(temp_file.name, 'rb'), as_attachment=True, filename=f'errors_batch_{batch.id}.xlsx')
        serializer = ImportErrorRowSerializer(rows, many=True)
        return Response(serializer.data)

    @action(detail=True, methods=['post'])
    def fix_error(self, request, project_id=None, pk=None):
        batch = self.get_object()
        row_id = request.data.get('row_id')
        resolved_data = request.data.get('resolved_data', {})

        try:
            error_row = ImportErrorRow.objects.get(id=row_id, import_batch=batch)
        except ImportErrorRow.DoesNotExist:
            return Response({'detail': '错误行不存在'}, status=status.HTTP_404_NOT_FOUND)

        error_row.resolved_data = resolved_data
        error_row.is_resolved = True
        error_row.save()

        return Response({'detail': '已修正', 'row_id': error_row.id})


class UnauditedTBViewSet(viewsets.ReadOnlyModelViewSet):
    serializer_class = UnauditedTBSerializer
    permission_classes = [IsProjectMember]

    def get_queryset(self):
        return UnauditedTB.objects.filter(
            project_id=self.kwargs['project_id']
        ).select_related('import_batch')

    @action(detail=False, methods=['get'])
    def current(self, request, project_id=None):
        rows = UnauditedTB.objects.filter(
            project_id=project_id,
            import_batch__is_active=True,
        ).select_related('import_batch')
        serializer = self.get_serializer(rows, many=True)
        return Response(serializer.data)


class UnauditedJEViewSet(viewsets.ReadOnlyModelViewSet):
    serializer_class = UnauditedJESerializer
    permission_classes = [IsProjectMember]

    def get_queryset(self):
        return UnauditedJE.objects.filter(
            project_id=self.kwargs['project_id']
        ).select_related('import_batch')

    @action(detail=False, methods=['get'])
    def current(self, request, project_id=None):
        rows = UnauditedJE.objects.filter(
            project_id=project_id,
            import_batch__is_active=True,
        ).select_related('import_batch')
        serializer = self.get_serializer(rows, many=True)
        return Response(serializer.data)


class ImportDashboardViewSet(viewsets.ViewSet):
    permission_classes = [IsProjectMember]
    pagination_class = StandardResultsSetPagination

    def list(self, request, project_id=None):
        queryset = ImportBatch.objects.filter(project_id=project_id, status='committed')
        period_filter = request.query_params.get('period')
        type_filter = request.query_params.get('import_type')
        if period_filter:
            queryset = queryset.filter(period=period_filter)
        if type_filter:
            queryset = queryset.filter(import_type=type_filter)

        # Group by period
        periods = sorted(set(q.period for q in queryset if q.period), reverse=True)
        result = []
        for p in periods:
            batches = queryset.filter(period=p).order_by('-version')
            result.append({
                'period': p,
                'batches': ImportBatchListSerializer(batches, many=True, context={'request': request}).data,
            })

        return Response(result)

    def retrieve(self, request, project_id=None, pk=None):
        try:
            batch = ImportBatch.objects.get(id=pk, project_id=project_id)
        except ImportBatch.DoesNotExist:
            return Response({'detail': '批次不存在'}, status=status.HTTP_404_NOT_FOUND)

        file_path = batch.metadata.get('file_path')
        mapping = batch.metadata.get('column_mapping') or batch.metadata.get('guessed_mapping')

        preview_rows = []
        if file_path and os.path.exists(file_path) and mapping:
            try:
                with open(file_path, 'rb') as f:
                    result = parse_file(f, batch.file_name)
                parsed = apply_mapping(result.rows[:200], mapping, batch.import_type)
                preview_rows = parsed['rows']
            except Exception:
                pass

        return Response({
            'batch': ImportBatchSerializer(batch).data,
            'preview_rows': preview_rows,
            'column_mapping': mapping,
        })

    @action(detail=True, methods=['get'])
    def download(self, request, project_id=None, pk=None):
        try:
            batch = ImportBatch.objects.get(id=pk, project_id=project_id)
        except ImportBatch.DoesNotExist:
            return Response({'detail': '批次不存在'}, status=status.HTTP_404_NOT_FOUND)

        file_path = batch.metadata.get('file_path')
        if not file_path or not os.path.exists(file_path):
            return Response({'detail': '文件不存在'}, status=status.HTTP_404_NOT_FOUND)

        return FileResponse(open(file_path, 'rb'), as_attachment=True, filename=batch.file_name)

    @action(detail=True, methods=['post'])
    def remap(self, request, project_id=None, pk=None):
        try:
            batch = ImportBatch.objects.get(id=pk, project_id=project_id)
        except ImportBatch.DoesNotExist:
            return Response({'detail': '批次不存在'}, status=status.HTTP_404_NOT_FOUND)

        batch.status = 'uploaded'
        batch.save(update_fields=['status'])

        ImportLog.objects.create(
            project=batch.project,
            import_batch=batch,
            action='remap',
            performed_by=request.user,
            details={'file_name': batch.file_name}
        )

        return Response({'batch_id': batch.id, 'status': batch.status})

    def destroy(self, request, project_id=None, pk=None):
        try:
            batch = ImportBatch.objects.get(id=pk, project_id=project_id)
        except ImportBatch.DoesNotExist:
            return Response({'detail': '批次不存在'}, status=status.HTTP_404_NOT_FOUND)

        ImportLog.objects.create(
            project=batch.project,
            import_batch=batch,
            action='delete',
            performed_by=request.user,
            details={'file_name': batch.file_name, 'period': batch.period}
        )

        file_path = batch.metadata.get('file_path')
        if file_path and os.path.exists(file_path):
            os.remove(file_path)

        batch.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


class MappingTemplateViewSet(viewsets.ModelViewSet):
    serializer_class = MappingTemplateSerializer
    permission_classes = [IsProjectMember]

    def get_queryset(self):
        return MappingTemplate.objects.filter(project_id=self.kwargs['project_id'])

    def perform_create(self, serializer):
        project = Project.objects.get(id=self.kwargs['project_id'])
        serializer.save(project=project)

    @action(detail=False, methods=['post'])
    def match(self, request, project_id=None):
        header_signature = request.data.get('header_signature', {})
        import_type = request.data.get('import_type', 'tb')

        templates = MappingTemplate.objects.filter(
            project_id=project_id,
            import_type=import_type,
        )

        best_match = None
        best_score = 0
        target_hash = header_signature.get('hash')
        target_count = header_signature.get('column_count', 0)

        for tmpl in templates:
            sig = tmpl.header_signature or {}
            score = 0
            if target_hash and sig.get('hash') == target_hash:
                score += 0.5
            col_count = sig.get('column_count', 0)
            if col_count and target_count:
                score += 0.3 * min(col_count, target_count) / max(col_count, target_count)
            if score > best_score:
                best_score = score
                best_match = tmpl

        if best_match and best_score >= 0.5:
            return Response({
                'matched': True,
                'score': round(best_score, 2),
                'template': MappingTemplateSerializer(best_match).data,
            })
        return Response({'matched': False, 'score': 0})


class ImportLogViewSet(viewsets.ReadOnlyModelViewSet):
    serializer_class = ImportLogSerializer
    permission_classes = [IsProjectMember]
    pagination_class = StandardResultsSetPagination

    def get_queryset(self):
        queryset = ImportLog.objects.filter(project_id=self.kwargs['project_id'])
        action_filter = self.request.query_params.get('action')
        if action_filter:
            queryset = queryset.filter(action=action_filter)
        return queryset.select_related('performed_by', 'import_batch')
